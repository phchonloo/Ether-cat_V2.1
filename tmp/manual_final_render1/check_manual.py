import hashlib
import json
from pathlib import Path

from docx import Document
from PIL import Image, ImageChops, ImageDraw, ImageFont


ROOT = Path(r"C:\Users\pc\Desktop\01_FCM4E1553_SQI_82_CTT_V1.4")
DOCX = ROOT / "docs" / "FCM4E1553_EtherCAT单轴步进电机驱动软件_使用说明书.docx"
REFERENCE = Path(r"C:\Users\pc\Desktop\组态屏软著\使用说明书.docx")
RENDER = ROOT / "tmp" / "manual_final_render4"
PAGE_DIR = RENDER / "pages"
pages = sorted(PAGE_DIR.glob("page-*.png"))

checks = []
for page_no, path in enumerate(pages, 1):
    with Image.open(path).convert("RGB") as image:
        white = Image.new("RGB", image.size, "white")
        diff = ImageChops.difference(image, white).convert("L").point(lambda x: 255 if x > 18 else 0)
        bbox = diff.getbbox()
        if bbox is None:
            checks.append({"page": page_no, "blank": True, "edge_risk": True, "bbox": None})
            continue
        left, top, right, bottom = bbox
        checks.append({
            "page": page_no,
            "blank": False,
            "edge_risk": left < 14 or top < 14 or right > image.width - 14 or bottom > image.height - 14,
            "bbox": list(bbox),
            "size": list(image.size),
        })

font = ImageFont.load_default()
contacts = []
for start in range(0, len(pages), 6):
    sheet = Image.new("RGB", (960, 1360), "#d6dbe3")
    draw = ImageDraw.Draw(sheet)
    for offset, path in enumerate(pages[start:start + 6]):
        page_no = start + offset + 1
        with Image.open(path).convert("RGB") as image:
            image.thumbnail((304, 620))
            col = offset % 3
            row = offset // 3
            x = col * 320 + (320 - image.width) // 2
            y = row * 680 + 24 + (620 - image.height) // 2
            sheet.paste(image, (x, y))
            draw.text((col * 320 + 8, row * 680 + 5), f"Page {page_no}", fill="black", font=font)
    out = RENDER / f"contact-{start + 1:02d}-{min(start + 6, len(pages)):02d}.png"
    sheet.save(out)
    contacts.append(str(out))

doc = Document(DOCX)
text = "\n".join(p.text for p in doc.paragraphs)
for table in doc.tables:
    for row in table.rows:
        text += "\n" + "\t".join(cell.text for cell in row.cells)

old_terms = [term for term in ("工业驱动调试组态屏", "Modbus RTU", "OTA", "授权码", "RTC 实时时钟") if term in text]
required_terms = [term for term in ("FCM4E1553", "EtherCAT", "CiA402", "PP", "PV", "CSP", "CSV", "6040", "6041", "SYNC0") if term in text]
report = {
    "docx": str(DOCX),
    "docx_sha256": hashlib.sha256(DOCX.read_bytes()).hexdigest(),
    "reference_unchanged": hashlib.sha256(REFERENCE.read_bytes()).hexdigest() == "ebc9c449ce989470c6f13c0d9a2ff1e14418244a04529993d8844430f648e87c",
    "pages": len(pages),
    "paragraphs": len(doc.paragraphs),
    "tables": len(doc.tables),
    "inline_shapes": len(doc.inline_shapes),
    "blank_pages": [x["page"] for x in checks if x["blank"]],
    "edge_risk_pages": [x["page"] for x in checks if x["edge_risk"]],
    "old_terms": old_terms,
    "required_terms_found": required_terms,
    "contacts": contacts,
}
(RENDER / "qa_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(report, ensure_ascii=False))
