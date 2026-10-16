from __future__ import annotations

from pathlib import Path
import hashlib
import json
import math
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor


ROOT = Path(r"C:\Users\pc\Desktop\01_FCM4E1553_SQI_82_CTT_V1.4")
SOURCE_ROOT = ROOT / "program" / "FCM4E1553 Ethercat"
OUTPUT = ROOT / "docs" / "FCM4E1553_源代码_60页.docx"
MANIFEST = ROOT / "tmp" / "source_doc_template" / "source_manifest.json"

SOURCE_EXTS = {".c", ".h", ".s", ".asm"}
PAGE_COUNT = 60
TARGET_LINES_PER_PAGE = 55

CODE_FONT = "Consolas"
CODE_SIZE = Pt(7)
CODE_LINE = Pt(7.3)

COLORS = {
    "plain": "3E4651",
    "comment": "8FA3B5",
    "preproc": "6B63D9",
    "keyword": "7057C7",
    "type": "3568D4",
    "string": "4E9E2A",
    "number": "D94C78",
    "function": "2E7BEF",
    "constant": "249AA6",
    "operator": "606770",
}

KEYWORDS = {
    "auto", "break", "case", "const", "continue", "default", "do", "else",
    "enum", "extern", "for", "goto", "if", "inline", "register", "restrict",
    "return", "sizeof", "static", "struct", "switch", "typedef", "union",
    "volatile", "while", "__asm", "__attribute__", "__inline", "__weak",
}
TYPES = {
    "bool", "char", "double", "float", "int", "long", "short", "signed",
    "unsigned", "void", "size_t", "ptrdiff_t", "uint8_t", "uint16_t",
    "uint32_t", "uint64_t", "int8_t", "int16_t", "int32_t", "int64_t",
    "u8", "u16", "u32", "s8", "s16", "s32", "UINT8", "UINT16", "UINT32",
    "UINT64", "INT8", "INT16", "INT32", "BOOL", "BYTE", "WORD", "DWORD",
}


def read_text(path: Path) -> str:
    data = path.read_bytes()
    for encoding in ("utf-8-sig", "gb18030", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            pass
    return data.decode("latin-1")


def source_lines(path: Path) -> list[str]:
    text = read_text(path).replace("\r\n", "\n").replace("\r", "\n")
    lines = text.split("\n")
    if lines and lines[-1] == "":
        lines.pop()
    return [clean_xml_text(line.expandtabs(4)) for line in lines]


def clean_xml_text(text: str) -> str:
    return "".join(ch for ch in text if ch in "\t\n\r" or ord(ch) >= 0x20)


def collect_sources() -> list[Path]:
    files = [
        p for p in SOURCE_ROOT.rglob("*")
        if p.is_file()
        and p.suffix.lower() in SOURCE_EXTS
        and "MDK-ARM/RTE/" not in p.relative_to(SOURCE_ROOT).as_posix()
    ]
    return sorted(files, key=lambda p: p.relative_to(SOURCE_ROOT).as_posix().lower())


def set_run_font(run, color: str, bold: bool = False) -> None:
    run.font.name = CODE_FONT
    run.font.size = CODE_SIZE
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(COLORS[color])
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), CODE_FONT)


def add_colored_run(paragraph, text: str, color: str, bold: bool = False) -> None:
    if not text:
        return
    run = paragraph.add_run(text)
    set_run_font(run, color, bold)


TOKEN_RE = re.compile(
    r"(//.*$|/\*|\*/|\"(?:\\.|[^\"\\])*\"|'(?:\\.|[^'\\])*'|"
    r"\b(?:0[xX][0-9A-Fa-f]+|\d+(?:\.\d+)?(?:[eE][+-]?\d+)?[uUlLfF]*)\b|"
    r"\b[A-Za-z_]\w*\b|==|!=|<=|>=|->|\+\+|--|&&|\|\||<<|>>|"
    r"[{}()\[\];,.?:~!%^&*+\-/=<>|])"
)


def style_identifier(token: str, line: str, end: int) -> tuple[str, bool]:
    if token in KEYWORDS:
        return "keyword", False
    if token in TYPES or token.endswith("_t"):
        return "type", False
    if token.isupper() and len(token) > 1:
        return "constant", False
    tail = line[end:]
    if re.match(r"\s*\(", tail):
        return "function", False
    return "plain", False


def highlight_line(line: str, in_block_comment: bool) -> tuple[list[tuple[str, str, bool]], bool]:
    if not line:
        return [], in_block_comment
    if in_block_comment:
        end = line.find("*/")
        if end < 0:
            return [(line, "comment", False)], True
        head = line[: end + 2]
        rest, state = highlight_line(line[end + 2 :], False)
        return [(head, "comment", False)] + rest, state
    if line.lstrip().startswith("#"):
        parts = []
        m = re.match(r"(\s*#\s*\w+)(.*)", line)
        if not m:
            return [(line, "preproc", False)], False
        parts.append((m.group(1), "preproc", False))
        remainder = m.group(2)
        if remainder:
            if "<" in remainder or '"' in remainder:
                parts.append((remainder, "string", False))
            else:
                parts.append((remainder, "plain", False))
        return parts, False

    result: list[tuple[str, str, bool]] = []
    pos = 0
    for match in TOKEN_RE.finditer(line):
        # A closed /* ... */ span may contain text that also matches TOKEN_RE.
        # Skip those already-consumed matches so comment contents are not emitted twice.
        if match.start() < pos:
            continue
        if match.start() > pos:
            result.append((line[pos : match.start()], "plain", False))
        token = match.group(0)
        if token.startswith("//"):
            result.append((token, "comment", False))
            pos = len(line)
            break
        if token == "/*":
            end = line.find("*/", match.end())
            if end < 0:
                result.append((line[match.start() :], "comment", False))
                return result, True
            result.append((line[match.start() : end + 2], "comment", False))
            pos = end + 2
            continue
        if token.startswith(('"', "'")):
            color, bold = "string", False
        elif re.fullmatch(r"0[xX][0-9A-Fa-f]+|\d+(?:\.\d+)?(?:[eE][+-]?\d+)?[uUlLfF]*", token):
            color, bold = "number", False
        elif re.fullmatch(r"[A-Za-z_]\w*", token):
            color, bold = style_identifier(token, line, match.end())
        else:
            color, bold = "operator", False
        result.append((token, color, bold))
        pos = match.end()
    if pos < len(line):
        result.append((line[pos:], "plain", False))
    return result, False


def set_paragraph_geometry(paragraph, page_break_before: bool = False) -> None:
    pf = paragraph.paragraph_format
    pf.left_indent = Pt(0)
    pf.right_indent = Pt(0)
    pf.first_line_indent = Pt(0)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = CODE_LINE
    pf.keep_together = True
    pf.keep_with_next = False
    pf.widow_control = False
    pf.page_break_before = page_break_before


def set_no_proof(run) -> None:
    rpr = run._element.get_or_add_rPr()
    no_proof = OxmlElement("w:noProof")
    rpr.append(no_proof)


def main() -> None:
    files = collect_sources()
    records = []
    all_lines: list[tuple[str, bool]] = []
    for path in files:
        lines = source_lines(path)
        records.append({
            "path": path.relative_to(SOURCE_ROOT).as_posix(),
            "lines": len(lines),
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        })
        in_block = False
        for line in lines:
            highlighted, in_block = highlight_line(line, in_block)
            all_lines.append((line, highlighted))

    if len(files) != 84 or len(all_lines) != 26803:
        raise RuntimeError(
            f"Source scope changed: files={len(files)}, lines={len(all_lines)}; expected 84/26803"
        )

    selected_line_count = PAGE_COUNT * TARGET_LINES_PER_PAGE
    selected_half = selected_line_count // 2
    document_lines = all_lines[:selected_half] + all_lines[-selected_half:]
    page_count = PAGE_COUNT
    page_sizes = [TARGET_LINES_PER_PAGE] * PAGE_COUNT
    break_indices = set()
    total = 0
    for size in page_sizes[:-1]:
        total += size
        break_indices.add(total)

    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(15)
    section.bottom_margin = Mm(15)
    section.left_margin = Mm(15)
    section.right_margin = Mm(15)
    section.header_distance = Mm(15)
    section.footer_distance = Mm(17.5)

    normal = doc.styles["Normal"]
    normal.font.name = CODE_FONT
    normal.font.size = CODE_SIZE
    normal._element.rPr.rFonts.set(qn("w:ascii"), CODE_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), CODE_FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), CODE_FONT)

    for index, (_plain, highlighted) in enumerate(document_lines):
        paragraph = doc.add_paragraph(style="Normal")
        set_paragraph_geometry(paragraph, page_break_before=index in break_indices)
        if not highlighted:
            run = paragraph.add_run("")
            set_run_font(run, "plain")
            set_no_proof(run)
            continue
        for text, color, bold in highlighted:
            run = paragraph.add_run(text)
            set_run_font(run, color, bold)
            set_no_proof(run)

    core = doc.core_properties
    core.title = "FCM4E1553 EtherCAT 单轴步进电机驱动软件源代码"
    core.subject = "软件著作权源代码前后各30页连续节选"
    core.author = ""
    core.last_modified_by = ""
    core.keywords = "FCM4E1553, EtherCAT, CiA402, N32G455"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)

    report = {
        "source_root": str(SOURCE_ROOT),
        "output": str(OUTPUT),
        "source_files": len(files),
        "source_lines": len(all_lines),
        "document_source_lines": len(document_lines),
        "selection": "first 1650 lines + last 1650 lines",
        "document_paragraphs": len(doc.paragraphs),
        "planned_pages": page_count,
        "logical_lines_per_page_min": min(page_sizes),
        "logical_lines_per_page_max": max(page_sizes),
        "files": records,
    }
    MANIFEST.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({k: v for k, v in report.items() if k != "files"}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
