import hashlib
import json
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(r"C:\Users\pc\Desktop\01_FCM4E1553_SQI_82_CTT_V1.4\tmp\manual_reference")
REFERENCE = Path(r"C:\Users\pc\Desktop\组态屏软著\使用说明书.docx")
PAGE_DIR = ROOT / "pages"

pages = sorted(PAGE_DIR.glob("page-*.png"))
font = ImageFont.load_default()
contacts = []
for start in range(0, len(pages), 6):
    sheet = Image.new("RGB", (960, 1360), "#d6dbe3")
    draw = ImageDraw.Draw(sheet)
    for offset, path in enumerate(pages[start:start + 6]):
        page_no = start + offset + 1
        with Image.open(path).convert("RGB") as image:
            image.thumbnail((304, 620))
            col = offset % 3
            row = offset // 3
            x = col * 320 + (320 - image.width) // 2
            y = row * 680 + 24 + (620 - image.height) // 2
            sheet.paste(image, (x, y))
            draw.text((col * 320 + 8, row * 680 + 5), f"Page {page_no}", fill="black", font=font)
    out = ROOT / f"contact-{start + 1:02d}-{min(start + 6, len(pages)):02d}.png"
    sheet.save(out)
    contacts.append(str(out))

doc = Document(REFERENCE)

def length_pt(value):
    return None if value is None else round(value.pt, 2)

def font_info(font):
    element = font._element
    rfonts = element.rPr.rFonts if element.rPr is not None else None
    return {
        "name": font.name,
        "east_asia": None if rfonts is None else rfonts.get(qn("w:eastAsia")),
        "size_pt": length_pt(font.size),
        "bold": font.bold,
        "italic": font.italic,
        "color": None if font.color is None or font.color.rgb is None else str(font.color.rgb),
    }

style_details = {}
for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Caption"):
    if name not in [style.name for style in doc.styles]:
        continue
    style = doc.styles[name]
    pf = style.paragraph_format
    style_details[name] = {
        "font": font_info(style.font),
        "alignment": None if pf.alignment is None else int(pf.alignment),
        "space_before_pt": length_pt(pf.space_before),
        "space_after_pt": length_pt(pf.space_after),
        "line_spacing": pf.line_spacing,
        "left_indent_pt": length_pt(pf.left_indent),
        "first_line_indent_pt": length_pt(pf.first_line_indent),
        "keep_with_next": pf.keep_with_next,
        "page_break_before": pf.page_break_before,
    }

section_details = []
for section in doc.sections:
    section_details.append({
        "page_width_pt": length_pt(section.page_width),
        "page_height_pt": length_pt(section.page_height),
        "top_margin_pt": length_pt(section.top_margin),
        "bottom_margin_pt": length_pt(section.bottom_margin),
        "left_margin_pt": length_pt(section.left_margin),
        "right_margin_pt": length_pt(section.right_margin),
        "header_distance_pt": length_pt(section.header_distance),
        "footer_distance_pt": length_pt(section.footer_distance),
        "header": [{"style": p.style.name, "text": p.text} for p in section.header.paragraphs],
        "footer": [{"style": p.style.name, "text": p.text} for p in section.footer.paragraphs],
    })
paragraphs = []
for i, paragraph in enumerate(doc.paragraphs):
    text = paragraph.text.strip()
    if text:
        paragraphs.append({"index": i, "style": paragraph.style.name, "text": text})

tables = []
for ti, table in enumerate(doc.tables):
    tables.append({
        "index": ti,
        "rows": len(table.rows),
        "cols": len(table.columns),
        "text": [[cell.text.strip() for cell in row.cells] for row in table.rows],
    })

with zipfile.ZipFile(REFERENCE) as archive:
    parts = {name: len(archive.read(name)) for name in archive.namelist()}

report = {
    "reference": str(REFERENCE),
    "sha256": hashlib.sha256(REFERENCE.read_bytes()).hexdigest(),
    "size": REFERENCE.stat().st_size,
    "rendered_pages": len(pages),
    "sections": len(doc.sections),
    "paragraph_count": len(doc.paragraphs),
    "table_count": len(doc.tables),
    "inline_shapes": len(doc.inline_shapes),
    "style_details": style_details,
    "section_details": section_details,
    "contacts": contacts,
    "paragraphs": paragraphs,
    "tables": tables,
    "parts": parts,
}
(ROOT / "reference_inventory.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps({k: v for k, v in report.items() if k not in {"paragraphs", "tables", "parts"}}, ensure_ascii=False))
